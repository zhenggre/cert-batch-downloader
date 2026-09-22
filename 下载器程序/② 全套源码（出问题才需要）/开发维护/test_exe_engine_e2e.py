# -*- coding: utf-8 -*-
"""对**打包后的 exe** 做引擎级端到端：验证「没有证书可下」这一整条链路。

为什么非得起个本地假网站：项目军规第 5 条 —— **验收对象是同事真正双击的那个 exe**。
源码版跑绿 ≠ 打包版没问题；而真实比赛网站当天抓不到（证书还没开放），
靠它反而验不了这一条。

做法是把 exe 当黑盒，只走它自己的 HTTP 接口：

    注入一个指向本地假网站的站点 → /api/start 跑 3 个人 → 读它生成的汇总清单

假网站对每个人的证书地址返回**空**（正是"证书还没挂出来"的样子）。于是必须满足：

  · 汇总清单「统计概览」里出现「暂无证书文件」= 3
  · 「下载失败」必须是 0（不许把"没开放"冒充成"失败"）
  · 界面上拿到的 result.nofile 必须 = 3、result.fail 必须 = 0
    ——否则同事对着一排 0，只会以为白跑了一趟
  · 不应该产出「失败名单」（它不是失败，不该进重跑队列）

D 段验的是**试跑**（2026-09-22 用户那次：筛了 604 人却只跑了 8 个，
报告上写着"全部完成"）：6 人的名单只抽查 2 人，必须从界面到日志到清单
到处写明"只查了 2 人、还有 4 人没跑、这不是全量"。

用法（CERTDL_RUN 指向 ① 里正在运行的 运行中.json）：
    set CERTDL_RUN=…\① …\证书下载器\运行中.json
    python test_exe_engine_e2e.py
"""
import base64
import json
import os
import shutil
import socket
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import requests

sys.stdout.reconfigure(encoding="utf-8")

RUN = Path(os.environ.get("CERTDL_RUN") or "")
if not RUN or not RUN.exists():
    print("没有拿到 CERTDL_RUN（要指向 ① 里正在运行的 运行中.json）")
    sys.exit(2)

run = json.loads(RUN.read_text(encoding="utf-8"))
PROG = RUN.parent
CFG = PROG / "config.json"
BASE = "http://127.0.0.1:%d" % run["port"]
TOK = run["token"]

ok = fail = 0


def check(name, cond, extra=""):
    global ok, fail
    if cond:
        ok += 1
        print("  [OK] %s" % name)
    else:
        fail += 1
        print("  [!!] %s   %s" % (name, extra))


# ---- 只跟本机说话：绝不能被沙箱/公司的代理劫持走 ----
S = requests.Session()
S.trust_env = False
URL_Q = "?token=" + TOK


# ---------------------------------------------------------------- 本地假网站
ROWS = {"F1": "端到端甲", "F999": "420100200801011234", "F2": "一等奖",
        "F3": "武汉市端到端小学", "F16": "湖北省", "F17": "武汉市",
        "F15": "洪山区", "F6": "王老师", "URL": ""}          # ← 证书地址故意留空


class Fake(BaseHTTPRequestHandler):
    def _send(self, obj, code=200):
        b = json.dumps(obj).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=UTF-8")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def do_GET(self):
        # 引擎建会话时会先 GET 一次 home_url 预热
        self._send({"ok": True})

    def do_POST(self):
        n = int(self.headers.get("Content-Length") or 0)
        self.rfile.read(n)                       # 请求体用不上，读掉就行
        # 永远"查得到这个人、但没有证书文件" —— 就是"成绩出了、证书还没挂出来"
        self._send({"total": 1, "rows": [ROWS]})

    def log_message(self, *a):
        pass


srv = ThreadingHTTPServer(("127.0.0.1", 0), Fake)
MPORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
print("本地假网站：http://127.0.0.1:%d" % MPORT)

backup = CFG.read_text(encoding="utf-8")
out_dir = Path(tempfile.mkdtemp(prefix="certdl_e2e_"))
SITE_KEY = "e2e_mock"

try:
    # ---------------------------------------------------------- 注入假站点
    cfg = json.loads(backup)
    cfg.setdefault("sites", {})[SITE_KEY] = {
        "title": "端到端测试（本地假网站）",
        "home_url": "http://127.0.0.1:%d/" % MPORT,
        "query_url": "http://127.0.0.1:%d/mock" % MPORT,
        "shortid": "e2e", "frmid": "e2e",
        "field_map": {"name": "F1", "id": "F999", "award": "F2", "school": "F3",
                      "province": "F16", "city": "F17", "district": "F15",
                      "teacher": "F6", "file_url": "URL"},
        "list_columns": {"name": ["姓名"], "id_number": ["证件号"]},
    }
    CFG.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")

    # ---------------------------------------------------------- 造一份小名单
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "证件号"])
    for i, (nm, idn) in enumerate([
            ("端到端甲", "420100200801011234"),
            ("端到端乙", "420100200802022345"),
            ("端到端丙", "420100200803033456")], 1):
        ws.append([nm, idn])
    # 注意：Windows 上 NamedTemporaryFile 会一直占着句柄，openpyxl 再往同一个
    # 路径保存会 PermissionError。所以先 mkstemp、马上关掉句柄、再写。
    fd, tmpname = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(tmpname)
    b64 = base64.b64encode(Path(tmpname).read_bytes()).decode()
    os.unlink(tmpname)

    # ---------------------------------------------------------- 起任务
    print("\n=== A. 让打包后的 exe 真跑一遍 ===")
    r = S.post(BASE + "/api/start" + URL_Q, json={
        "filename": "端到端名单.xlsx", "content_b64": b64,
        "out": str(out_dir), "subfolder": False,
        "site": SITE_KEY, "sample": 0, "filters": {},
    }, timeout=30)
    j = r.json()
    check("接口接受了任务", j.get("ok") is True, json.dumps(j)[:200])

    deadline = time.time() + 120
    st = {}
    while time.time() < deadline:
        st = S.get(BASE + "/api/progress" + URL_Q, timeout=30).json()
        if not st.get("running"):
            break
        time.sleep(1)
    check("任务跑完了（没卡死）", not st.get("running"), st.get("stage"))
    check("过程里没有异常", not st.get("error"), str(st.get("error"))[:200])

    res = st.get("result") or {}
    print("     界面拿到的结果：%s" % json.dumps(
        {k: res.get(k) for k in ("cert", "schools", "miss", "fail", "nofile")},
        ensure_ascii=False))

    print("\n=== B. 界面上必须能看出来「不是失败」 ===")
    check("result 里有 nofile 这个字段（界面要显示它）", "nofile" in res,
          "keys=%s" % sorted(res.keys()))
    check("nofile = 3（三个人都是「没有证书可下」）", res.get("nofile") == 3,
          "实际=%r" % res.get("nofile"))
    check("fail = 0（没开放 ≠ 失败）", res.get("fail") == 0,
          "实际=%r" % res.get("fail"))
    check("cert = 0（确实一份都没有）", res.get("cert") == 0,
          "实际=%r" % res.get("cert"))
    check("没生成失败名单（不该进重跑队列）", not res.get("retry_xlsx"),
          str(res.get("retry_xlsx")))

    print("\n=== C. 汇总清单里要把话说清楚 ===")
    xlsx = res.get("xlsx") or ""
    check("拿到了汇总清单路径", bool(xlsx) and Path(xlsx).exists(), xlsx)
    if xlsx and Path(xlsx).exists():
        ws4 = load_workbook(xlsx)["统计概览"]
        rows = {}
        for row in ws4.iter_rows(values_only=True):
            if row and row[0]:
                rows[str(row[0])] = row[1]
        lab = [k for k in rows if str(k).startswith("暂无证书文件")]
        check("统计概览里有「暂无证书文件」这一行", bool(lab),
              "只有 %s" % list(rows))
        if lab:
            check("这一行数对了（3 人）", rows[lab[0]] == 3,
                  "实际=%r" % rows[lab[0]])
        check("「下载失败」是 0", rows.get("下载失败") == 0,
              "实际=%r" % rows.get("下载失败"))

    # ---------------------------------------------------------- D. 试跑模式
    # 2026-09-22 的真实事故：筛出 604 人、「先随机试跑」是默认勾着的，于是只跑了
    # 8 个人，而报告上写着「全部完成 ✅」、清单里写着「名单总人数 8」。
    # 这一节就让打包后的 exe 真跑一次"6 人的名单、只抽查 2 人"，
    # 看它有没有从日志到清单都老老实实交代"这不是全量"。
    print("\n=== D. 试跑：6 人的名单只抽查 2 人，必须到处写明「不是全量」===")
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["姓名", "证件号"])
    for i in range(6):
        ws2.append(["端到端第%d人" % (i + 1), "42010020080101%04d" % (1000 + i)])
    fd2, tmp2 = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd2)
    wb2.save(tmp2)
    b64b = base64.b64encode(Path(tmp2).read_bytes()).decode()
    os.unlink(tmp2)

    out2 = Path(tempfile.mkdtemp(prefix="certdl_trial_"))
    S.post(BASE + "/api/start" + URL_Q, json={
        "filename": "试跑名单.xlsx", "content_b64": b64b,
        "out": str(out2), "subfolder": False,
        "site": SITE_KEY, "sample": 2, "filters": {},
    }, timeout=30)
    saw_trial, st2, deadline = False, {}, time.time() + 120
    while time.time() < deadline:
        st2 = S.get(BASE + "/api/progress" + URL_Q, timeout=30).json()
        saw_trial = saw_trial or bool(st2.get("trial"))
        if not st2.get("running"):
            break
        time.sleep(0.5)
    res2 = st2.get("result") or {}
    check("跑的过程中界面能收到「试跑」标记", saw_trial)
    check("result.sampled = 2（本次只查了 2 人）", res2.get("sampled") == 2,
          "实际=%r" % res2.get("sampled"))
    check("result.filtered_total = 6（筛完其实有 6 人，不是 2 人）",
          res2.get("filtered_total") == 6, "实际=%r" % res2.get("filtered_total"))

    logs = sorted(out2.glob("运行日志_*.txt"))
    check("生成了运行日志", bool(logs), str(out2))
    if logs:
        txt = logs[-1].read_text(encoding="utf-8-sig")
        check("日志开头喊了「这是试跑」", "★ 这是「试跑」" in txt)
        check("日志写明筛完 6 人、本次只查 2 人",
              "筛选后名单有 6 人" in txt and "只随机查了 2 人" in txt)
        check("日志结尾写的是「试跑结束」", "试跑结束" in txt)
        # 最要紧的一条：「全部完成」这四个字绝不能出现在试跑里。
        check("日志里没有「全部完成」这种误导说法", "全部完成" not in txt,
              "出现了：%s" % [l for l in txt.splitlines() if "全部完成" in l][:2])

    x2 = res2.get("xlsx") or ""
    check("拿到了试跑的汇总清单", bool(x2) and Path(x2).exists(), x2)
    if x2 and Path(x2).exists():
        rows2 = {}
        for row in load_workbook(x2)["统计概览"].iter_rows(values_only=True):
            if row and row[0]:
                rows2[str(row[0])] = row[1]
        head = list(rows2)[:4]
        check("「试跑」写在了统计概览的最前面（只截第一屏也看得到）",
              any(str(k).startswith("⚠ 本次是「试跑」") for k in head),
              "前四行=%s" % head)
        check("清单写明「筛选后名单总人数（本次没跑）」= 6",
              rows2.get("   筛选后名单总人数（本次没跑）") == 6,
              "实际=%r" % rows2.get("   筛选后名单总人数（本次没跑）"))
        check("清单写明「本次实际查询人数」= 2",
              rows2.get("   本次实际查询人数") == 2,
              "实际=%r" % rows2.get("   本次实际查询人数"))
    shutil.rmtree(out2, ignore_errors=True)

finally:
    CFG.write_text(backup, encoding="utf-8")
    print("\n（config.json 已还原成测试前的样子）")
    srv.shutdown()
    shutil.rmtree(out_dir, ignore_errors=True)

print()
print("=" * 66)
print("结果：%d 项通过，%d 项失败" % (ok, fail))
print("=" * 66)
sys.exit(1 if fail else 0)
