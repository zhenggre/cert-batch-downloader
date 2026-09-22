# -*- coding: utf-8 -*-
"""新状态「暂无证书文件」的独立检测点（2026-09-21 用户那次国赛抓出来的）。

背景：比赛成绩刚出、证书还没开放下载时，页面上每个人取到的证书地址都是**空**。
老代码把它当「下载失败」报，还抛 MissingSchema 这种同事看不懂的东西 ——
既是误导（用户当场就问"是不是他们没开放下载？"），又白喂退避/熔断计数，
跑出一份全是假失败的名单。

这条属于典型的"不报错的错"：它不崩、不抛、不亮红灯，只是把结果分错类。
所以必须有人专门守着，不能指望它自己冒头。

盯三件事：
  A. 判定规则本身（纯函数真值表）—— 空地址=暂无；有地址=该去下；续跑命中=跳过优先
  B. 接线：没地址时**绝不发请求**（否则又会拿 MissingSchema 冒充"网络失败"）
  C. 汇总清单里真的多出这一行，且**不混进「下载失败」**
"""
import shutil
import sys
from pathlib import Path

R = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(R / "程序"))
import cert_downloader as core      # noqa: E402

ok = fail = 0


def check(name, cond, extra=""):
    global ok, fail
    if cond:
        ok += 1
        print("  [OK] %s" % name)
    else:
        fail += 1
        print("  [!!] %s   %s" % (name, extra))


TMP = Path(__file__).resolve().parent / "_tmp_nofile_test"
if TMP.exists():
    shutil.rmtree(TMP)
TMP.mkdir(parents=True)

print("=" * 66)
print("A. 判定规则（不看网络，只看有没有地址）")
print("=" * 66)
f = core.is_no_file
check("地址是空串 → 暂无证书文件", f(None, "") is True)
check("地址是纯空白 → 暂无证书文件", f(None, "   ") is True)
check("地址是 None → 暂无证书文件", f(None, None) is True)
check("地址有值 → 不是这一类（该去下）", f(None, "https://x/a.pdf") is False)
check("续跑已命中老文件（target 有值）→ 跳过优先，不算暂无",
      f(Path("x.pdf"), "") is False)
check("返回真布尔值（能被 if 直接用，不是真值对象）",
      isinstance(f(None, ""), bool))

print()
print("=" * 66)
print("B. 接线：没地址时不许去发请求")
print("=" * 66)
src = (R / "程序" / "cert_downloader.py").read_text(encoding="utf-8")
check("no_file 由 is_no_file() 算出来（没被就地改回去）",
      'no_file = is_no_file(target, rec.get("file_url"))' in src)
check("下载分支带 not no_file 护栏（没地址就不下）",
      "if target is None and not no_file:" in src)
i_guard = src.find("if target is None and not no_file:")
i_fetch = src.find("body, err, tries = fetch_pdf(")
check("护栏在 fetch_pdf 之前（顺序对了才拦得住）",
      i_guard != -1 and i_fetch != -1 and i_guard < i_fetch,
      "guard=%d fetch=%d" % (i_guard, i_fetch))
# 同一件事还有另一半来源：站点**根本没有证书文件字段**（金数据常见）。
# 它原先写的是另一个状态 "无证书文件"，而那一处统计/日志/界面/测试四处都没接，
# 等于隐形 —— 同事看到的还是一排 0。现在两半并成同一个状态，靠 note 区分原因。
check("站点没证书字段时也归到同一个状态（不再另起隐形状态）",
      "查询页没有证书文件字段（主办方没把证书放进可查询结果）" in src)
# 判据盯"赋值"，不盯"出现"：注释里为了讲清历史会引用旧名，那是应该的。
check('不再有「无证书文件」这个状态被赋给记录（旧名已并入）',
      '"status": "无证书文件"' not in src
      and 'status = "无证书文件"' not in src)

print()
print("=" * 66)
print("C. 汇总清单：真的多一行，且不混进「下载失败」")
print("=" * 66)


def rec(name, idn, school, status, note=""):
    return {"name": name, "id": idn, "school": school, "raw_school": school,
            "official_school": school,
            "award": "一等奖" if status == "成功" else "",
            "province": "湖北省", "city": "武汉市", "district": "洪山区",
            "teacher": "王老师", "file_url": "", "hit": status == "成功",
            "qerr": "", "total": 1 if status == "成功" else 0, "status": status,
            "saved_path": "", "saved_kb": 0, "note": note,
            "cert_ok": status == "成功"}


NOTE = "比赛还没开放下载（页面上还没有证书文件）"
NOTE_SITE = "查询页没有证书文件字段（主办方没把证书放进可查询结果）"
records = [
    rec("陈一", "420100200801011234", "武汉市洪山区第一小学", "暂无证书文件", NOTE),
    rec("陈二", "420100200802022345", "武汉市洪山区第一小学", "暂无证书文件", NOTE),
    rec("陈三", "420100200803033456", "武汉市洪山区第二小学", "暂无证书文件", NOTE_SITE),
    rec("陈四", "420100200804044567", "武汉市洪山区第二小学", "下载失败",
        "等待 120s 仍未生成完成"),
    rec("陈五", "420100200805055678", "武汉市洪山区第二小学", "未获奖无证书"),
]
out = TMP / "清单"
out.mkdir(parents=True)
xlsx = out / "证书汇总清单_测试.xlsx"
core.write_summary(xlsx, records, out, disk={})

from openpyxl import load_workbook        # noqa: E402
ws = load_workbook(xlsx)["统计概览"]
rows = {}
for r in ws.iter_rows(values_only=True):
    if r and r[0]:
        rows[str(r[0])] = r[1]

label = "暂无证书文件（比赛还没开放下载 / 查询页未提供）"
check("统计概览里确有这一行", label in rows, "只有 %s" % list(rows))
check("这一行把两种来源都数进去了（3 人）", rows.get(label) == 3,
      "实际=%r" % rows.get(label))
check("「下载失败」没把暂无算进去（仍是 1）", rows.get("下载失败") == 1,
      "实际=%r" % rows.get("下载失败"))
check("两个口径分开，不会让人以为整批都失败",
      rows.get(label) == 3 and rows.get("下载失败") == 1)

# 顺便确认这一条在明细里也说得清楚（同事看不懂等于没说）
ws2 = load_workbook(xlsx)["未获奖及异常"]
text = "\n".join(" ".join(str(c) for c in r if c)
                 for r in ws2.iter_rows(values_only=True))
check("明细里写了「还没开放下载」这类人话",
      "还没开放下载" in text, text[:120].replace("\n", " / "))
check("站点没证书字段那一种也把原因写清了",
      "没有证书文件字段" in text, text[:200].replace("\n", " / "))

shutil.rmtree(TMP, ignore_errors=True)

print()
print("=" * 66)
print("结果：%d 项通过，%d 项失败" % (ok, fail))
print("=" * 66)
sys.exit(1 if fail else 0)
